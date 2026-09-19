"""Readers for charts of accounts and transaction journals.

Accepts both CSV and .xlsx so the same code path handles the supplied QuickBooks
export and the SaaS transaction file. Column headers are matched loosely, since
every accounting export names things slightly differently.
"""

from __future__ import annotations

import csv
import io
from pathlib import Path

from openpyxl import load_workbook

from .models import Account, TransactionLine

# Header aliases -> canonical field name. Compared case-insensitively with
# punctuation stripped, so "Account Name" and "account_name" both resolve.
_ACCOUNT_ALIASES = {
    "accountnumber": "account_number",
    "account": "account_number",
    "number": "account_number",
    "acct": "account_number",
    "accountname": "account_name",
    "name": "account_name",
    "description": "account_name",
    "type": "type",
    "accounttype": "type",
    "detailtype": "detail_type",
    "detail": "detail_type",
    "subtype": "detail_type",
}

_TRANSACTION_ALIASES = {
    "transactionid": "transaction_id",
    "transid": "transaction_id",
    "id": "transaction_id",
    "date": "date",
    "transactiondate": "date",
    "accountnumber": "account_number",
    "account": "account_number",
    "accountname": "account_name",
    "amount": "amount",
    "class": "cls",
    "classname": "cls",
    "location": "location",
    "memo": "memo",
    "description": "memo",
}


class ParseError(ValueError):
    """Raised when a file has no usable header row."""


def _canon_header(value: object) -> str:
    return "".join(ch for ch in str(value or "").lower() if ch.isalnum())


def _map_headers(raw: list[object], aliases: dict[str, str]) -> dict[int, str]:
    """Return {column index -> canonical field} for the headers we recognise."""
    resolved: dict[int, str] = {}
    for idx, cell in enumerate(raw):
        canonical = aliases.get(_canon_header(cell))
        # First column wins if an export repeats a header.
        if canonical and canonical not in resolved.values():
            resolved[idx] = canonical
    return resolved


def _rows_from_bytes(data: bytes, filename: str) -> list[list[object]]:
    """Read a CSV or XLSX payload into a list of rows."""
    if Path(filename).suffix.lower() in {".xlsx", ".xlsm"}:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.active
        return [list(row) for row in sheet.iter_rows(values_only=True)]

    # utf-8-sig strips the BOM that Excel writes on "Save as CSV".
    text = data.decode("utf-8-sig", errors="replace")
    return [list(row) for row in csv.reader(io.StringIO(text))]


def _to_float(value: object) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    cleaned = str(value).replace(",", "").replace("$", "").strip()
    # Parenthesised negatives, as produced by most accounting exports.
    if cleaned.startswith("(") and cleaned.endswith(")"):
        cleaned = "-" + cleaned[1:-1]
    try:
        return float(cleaned)
    except ValueError:
        return 0.0


def parse_accounts(data: bytes, filename: str) -> list[Account]:
    """Parse a chart of accounts. Blank rows are skipped."""
    rows = _rows_from_bytes(data, filename)
    if not rows:
        raise ParseError("The file is empty.")

    header = _map_headers(rows[0], _ACCOUNT_ALIASES)
    if "account_number" not in header.values():
        raise ParseError(
            "No account-number column found. Expected a header like "
            "'account_number' or 'Account'."
        )

    accounts: list[Account] = []
    for row in rows[1:]:
        values: dict[str, object] = {}
        for idx, field in header.items():
            values[field] = row[idx] if idx < len(row) else None

        number = str(values.get("account_number") or "").strip()
        name = str(values.get("account_name") or "").strip()
        if not number and not name:
            continue
        if not number:
            # Keep the row but flag it by name only; the matcher tolerates this.
            number = ""
        accounts.append(
            Account(
                account_number=number,
                account_name=name or number,
                type=str(values.get("type") or "").strip(),
                detail_type=str(values.get("detail_type") or "").strip(),
            )
        )
    return accounts


def parse_transactions(data: bytes, filename: str) -> list[TransactionLine]:
    """Parse a transaction journal. Rows without an amount are skipped."""
    rows = _rows_from_bytes(data, filename)
    if not rows:
        raise ParseError("The file is empty.")

    header = _map_headers(rows[0], _TRANSACTION_ALIASES)
    if "account_number" not in header.values():
        raise ParseError(
            "No account-number column found in the transaction file."
        )

    lines: list[TransactionLine] = []
    for row in rows[1:]:
        values: dict[str, object] = {}
        for idx, field in header.items():
            values[field] = row[idx] if idx < len(row) else None

        number = str(values.get("account_number") or "").strip()
        amount_raw = values.get("amount")
        if not number or amount_raw in (None, ""):
            continue

        date_value = values.get("date")
        if hasattr(date_value, "strftime"):
            date_text = date_value.strftime("%Y-%m-%d")
        else:
            date_text = str(date_value or "").strip()

        lines.append(
            TransactionLine(
                transaction_id=str(values.get("transaction_id") or "").strip(),
                date=date_text,
                account_number=number,
                account_name=str(values.get("account_name") or "").strip(),
                amount=_to_float(amount_raw),
                cls=str(values.get("cls") or "").strip(),
                location=str(values.get("location") or "").strip(),
                memo=str(values.get("memo") or "").strip(),
            )
        )
    return lines
